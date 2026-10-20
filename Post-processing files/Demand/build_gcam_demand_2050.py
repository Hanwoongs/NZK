#!/usr/bin/env python3
"""
Rebuild demand_ref.csv / demand_NDC.csv (2021-2050, provincial, 8760h) using
the new GCAM annual electricity-demand totals in ../../input.xlsx.

Method
  For 2021/2025/2030/2035: take that year's own hourly, 17-province tab in
  regional_hourly_load_profile_2024_analysis.xlsx (already internally
  consistent: sum(provinces) = national, every hour). Uniformly rescale
  every value in that tab (national + all 17 provinces) by
  scale = target_annual_MWh / original_tab_annual_MWh
  so the hourly shape AND the provincial split are both preserved exactly;
  only the overall level is updated to match the new GCAM total.

  For 2040/2045/2050 (no tab exists yet): reuse the 2035 tab as the shape
  template and rescale it the same way, using each year's own GCAM target.

  target_annual_MWh comes from input.xlsx "Elec demand" (EJ) x 277,777,778
  (MWh/EJ), consistent with the conversion already used for this project's
  GCAM-MACRO shared-input workbook.

Output
  runs/Ref/system/demand_ref.csv   (elec cols replaced, heat/Zero kept)
  runs/NDC/system/demand_NDC.csv   (elec cols replaced, heat/Zero kept)
"""
from __future__ import annotations
import csv
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent.parent  # .../NZK_power provincial_0908_BNEF
XLSX_PROFILE = (ROOT / "Post-processing files/Demand"
                "/National_to_Regional_Downscaling_script_py_version"
                "/Input_data/regional_hourly_load_profile_2024_analysis.xlsx")
XLSX_GCAM = ROOT / "input.xlsx"
EJ_TO_MWH = 277_777_778.0

NAME2CODE = {
    "Gangwon": "GWN", "Gyeonggi": "GGI", "Gyeongnam": "GNA", "Gyeongbuk": "GNB",
    "Gwangju": "KWJ", "Daegu": "TAE", "Daejeon": "DJJ", "Busan": "PUS",
    "Seoul": "SEL", "Sejong": "SJG", "Ulsan": "USN", "Incheon": "INC",
    "Jeonnam": "JNA", "Jeonbuk": "JNB", "Jeju": "JEJ", "Chungnam": "CNA",
    "Chungbuk": "CNB",
}
YEAR_SUFFIX = {2021: "", 2025: "_25", 2030: "_30", 2035: "_35",
               2040: "_40", 2045: "_45", 2050: "_50"}
TAB_YEARS = (2021, 2025, 2030, 2035)  # years that have their own workbook tab


def read_targets(scenario: str) -> dict[int, float]:
    """year -> target annual MWh, from input.xlsx."""
    wb = openpyxl.load_workbook(XLSX_GCAM, read_only=True, data_only=True)
    sheet = "ref" if scenario == "ref" else "NDC"
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    years = [c for c in hdr if isinstance(c, int)]
    demand_row = next(r for r in rows[1:] if r[0] == "Elec demand")
    out = {}
    for yr in years:
        j = hdr.index(yr)
        out[yr] = float(demand_row[j]) * EJ_TO_MWH
    return out


def read_tab(wb, tab_name: str):
    """-> (national[8760] list, {code: series[8760]}) from one workbook tab."""
    ws = wb[tab_name]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    prov_col = {}
    for j, h in enumerate(hdr):
        if h in NAME2CODE:
            prov_col[NAME2CODE[h]] = j
    nat, prov = [], {c: [] for c in prov_col}
    for r in rows[1:]:
        if r[1] is None:
            continue
        nat.append(float(r[1]))
        for c, j in prov_col.items():
            prov[c].append(float(r[j]))
    assert len(nat) == 8760, f"{tab_name}: got {len(nat)} rows, expected 8760"
    assert len(prov) == 17, f"{tab_name}: got {len(prov)} provinces, expected 17"
    return nat, prov


def build(scenario: str):
    targets = read_targets(scenario)
    print(f"\n=== {scenario} targets (MWh) ===")
    for yr, v in targets.items():
        print(f"  {yr}: {v/1e6:8.2f} TWh")

    wb = openpyxl.load_workbook(XLSX_PROFILE, read_only=True, data_only=True)

    result = {}  # year -> (nat[8760], {code: series[8760]})
    tab_cache = {}
    for yr in YEAR_SUFFIX:
        src_year = yr if yr in TAB_YEARS else 2035
        if src_year not in tab_cache:
            tab_cache[src_year] = read_tab(wb, f"{src_year}_{scenario}")
        nat0, prov0 = tab_cache[src_year]
        scale = targets[yr] / sum(nat0)
        nat = [v * scale for v in nat0]
        prov = {c: [v * scale for v in s] for c, s in prov0.items()}
        result[yr] = (nat, prov)
        tag = "" if src_year == yr else f" (shape from {src_year})"
        print(f"  {yr}: scale={scale:.4f}{tag}")

    return result


def write_csv(scenario: str, result: dict, template_path: Path, out_path: Path):
    with open(template_path) as fh:
        rd = csv.DictReader(fh)
        template_hdr = rd.fieldnames
        template_rows = list(rd)
    assert len(template_rows) == 8760

    out_rows = []
    for i in range(8760):
        row = dict(template_rows[i])  # carries Demand_heat*, Demand_Zero as-is
        row["Time_Index"] = i + 1
        for yr, suf in YEAR_SUFFIX.items():
            nat, prov = result[yr]
            row[f"Demand_MW{suf}"] = round(nat[i], 6)
            for code, series in prov.items():
                row[f"Demand_MW{suf}_{code}"] = round(series[i], 6)
        out_rows.append(row)

    with open(out_path, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=template_hdr)
        w.writeheader()
        w.writerows(out_rows)
    print(f"wrote {out_path}")


def main():
    ref = build("ref")
    ndc = build("NDC")

    write_csv("ref", ref,
              ROOT / "runs/Ref/system/demand_ref.csv",
              ROOT / "runs/Ref/system/demand_ref.csv")
    write_csv("NDC", ndc,
              ROOT / "runs/NDC/system/demand_NDC.csv",
              ROOT / "runs/NDC/system/demand_NDC.csv")


if __name__ == "__main__":
    main()
